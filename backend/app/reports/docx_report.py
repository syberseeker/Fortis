from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .schema import SecurityReport, SEVERITY_ORDER

SEVERITY_COLORS = {
    "Critical": RGBColor(0x8B, 0x00, 0x00),
    "High": RGBColor(0xC0, 0x39, 0x2B),
    "Medium": RGBColor(0xE6, 0x7E, 0x22),
    "Low": RGBColor(0x21, 0x87, 0x38),
    "Informational": RGBColor(0x5D, 0x6D, 0x7E),
}


def render_docx(report: SecurityReport, output_path: str) -> str:
    doc = Document()

    # Title page
    title = doc.add_heading(report.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("Cybersecurity Assessment Report")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor(0x5D, 0x6D, 0x7E)
    date_p = doc.add_paragraph(datetime.now().strftime("%d %B %Y"))
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    doc.add_heading("1. Client Context", level=1)
    doc.add_paragraph(report.client_context)

    doc.add_heading("2. Scope", level=1)
    doc.add_paragraph(report.scope)

    doc.add_heading("3. Executive Summary", level=1)
    doc.add_paragraph(report.executive_summary)
    p = doc.add_paragraph()
    p.add_run("Overall Risk Rating: ").bold = True
    run = p.add_run(report.overall_risk_rating)
    run.bold = True
    run.font.color.rgb = SEVERITY_COLORS.get(report.overall_risk_rating, RGBColor(0, 0, 0))

    doc.add_heading("4. Findings", level=1)
    sorted_findings = sorted(
        report.findings,
        key=lambda f: SEVERITY_ORDER.index(f.severity) if f.severity in SEVERITY_ORDER else 99,
    )
    if not sorted_findings:
        doc.add_paragraph("No findings were identified in the reviewed material.")

    for i, finding in enumerate(sorted_findings, start=1):
        h = doc.add_heading(f"4.{i} {finding.title}", level=2)
        sev_p = doc.add_paragraph()
        sev_p.add_run("Severity: ").bold = True
        sev_run = sev_p.add_run(finding.severity)
        sev_run.bold = True
        sev_run.font.color.rgb = SEVERITY_COLORS.get(finding.severity, RGBColor(0, 0, 0))

        if finding.source_file:
            src_p = doc.add_paragraph()
            src_p.add_run("Source: ").bold = True
            src_p.add_run(finding.source_file)

        if finding.framework and finding.control_id:
            fw_p = doc.add_paragraph()
            fw_p.add_run("Framework Mapping: ").bold = True
            fw_p.add_run(f"{finding.framework} — {finding.control_id}")

        doc.add_paragraph("Description:", style="Intense Quote")
        doc.add_paragraph(finding.description)

        doc.add_paragraph("Evidence:", style="Intense Quote")
        doc.add_paragraph(finding.evidence)

        doc.add_paragraph("Remediation:", style="Intense Quote")
        doc.add_paragraph(finding.remediation)

        doc.add_paragraph()  # spacing

    doc.add_heading("5. Recommendations Summary", level=1)
    for rec in report.recommendations_summary:
        doc.add_paragraph(rec, style="List Bullet")

    doc.save(output_path)
    return output_path
