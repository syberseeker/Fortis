from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from .schema import SecurityReport, SEVERITY_ORDER

SEVERITY_COLORS = {
    "Critical": RGBColor(0x8B, 0x00, 0x00),
    "High": RGBColor(0xC0, 0x39, 0x2B),
    "Medium": RGBColor(0xE6, 0x7E, 0x22),
    "Low": RGBColor(0x21, 0x87, 0x38),
    "Informational": RGBColor(0x5D, 0x6D, 0x7E),
}

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
SLATE = RGBColor(0x5D, 0x6D, 0x7E)


def _add_title_slide(prs, report: SecurityReport):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = report.title
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(36)
    subtitle = slide.placeholders[1]
    subtitle.text = f"Cybersecurity Assessment — {datetime.now().strftime('%d %B %Y')}"
    return slide


def _add_bullet_slide(prs, heading: str, bullets: list, sub=None):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = heading
    body = slide.placeholders[1].text_frame
    body.clear()
    first = True
    if sub:
        body.text = sub
        first = False
    for b in bullets:
        p = body.paragraphs[0] if first else body.add_paragraph()
        p.text = b
        p.level = 0
        first = False
    return slide


def _add_finding_slide(prs, idx: int, finding):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = f"Finding {idx}: {finding.title}"

    body = slide.placeholders[1].text_frame
    body.clear()

    p = body.paragraphs[0]
    run = p.add_run()
    run.text = f"Severity: {finding.severity}"
    run.font.bold = True
    run.font.color.rgb = SEVERITY_COLORS.get(finding.severity, RGBColor(0, 0, 0))

    if finding.source_file:
        p2 = body.add_paragraph()
        p2.text = f"Source: {finding.source_file}"

    if finding.framework and finding.control_id:
        p3 = body.add_paragraph()
        p3.text = f"Mapped control: {finding.framework} — {finding.control_id}"

    p4 = body.add_paragraph()
    p4.text = f"Description: {finding.description}"

    p5 = body.add_paragraph()
    p5.text = f"Remediation: {finding.remediation}"

    return slide


def render_pptx(report: SecurityReport, output_path: str) -> str:
    prs = Presentation()

    _add_title_slide(prs, report)

    _add_bullet_slide(
        prs,
        "Executive Summary",
        [report.executive_summary, f"Overall Risk Rating: {report.overall_risk_rating}"],
    )

    _add_bullet_slide(prs, "Scope & Context", [report.client_context, report.scope])

    sorted_findings = sorted(
        report.findings,
        key=lambda f: SEVERITY_ORDER.index(f.severity) if f.severity in SEVERITY_ORDER else 99,
    )

    if not sorted_findings:
        _add_bullet_slide(prs, "Findings", ["No findings identified in the reviewed material."])
    else:
        severity_counts = {}
        for f in sorted_findings:
            severity_counts[f.severity] = severity_counts.get(f.severity, 0) + 1
        overview_bullets = [f"{sev}: {count}" for sev, count in severity_counts.items()]
        _add_bullet_slide(prs, "Findings Overview", overview_bullets)

        for i, finding in enumerate(sorted_findings, start=1):
            _add_finding_slide(prs, i, finding)

    _add_bullet_slide(prs, "Recommendations", report.recommendations_summary)

    prs.save(output_path)
    return output_path
