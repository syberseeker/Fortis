import tempfile
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .schema import SecurityReport, SEVERITY_ORDER
from .charts import render_severity_chart, render_framework_chart
from ..structured import is_valid_mermaid

SEVERITY_COLORS = {
    "Critical": RGBColor(0x8B, 0x00, 0x00),
    "High": RGBColor(0xC0, 0x39, 0x2B),
    "Medium": RGBColor(0xE6, 0x7E, 0x22),
    "Low": RGBColor(0x21, 0x87, 0x38),
    "Informational": RGBColor(0x5D, 0x6D, 0x7E),
}


def render_docx(report: SecurityReport, output_path: str) -> str:
    doc = Document()

    # Title page
    title = doc.add_heading(report.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("Cybersecurity Assessment Report")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor(0x5D, 0x6D, 0x7E)
    date_p = doc.add_paragraph(datetime.now().strftime("%d %B %Y"))
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    doc.add_heading("1. Client Context", level=1)
    doc.add_paragraph(report.client_context)

    doc.add_heading("2. Scope", level=1)
    doc.add_paragraph(report.scope)

    doc.add_heading("3. Executive Summary", level=1)
    doc.add_paragraph(report.executive_summary)
    p = doc.add_paragraph()
    p.add_run("Overall Risk Rating: ").bold = True
    run = p.add_run(report.overall_risk_rating)
    run.bold = True
    run.font.color.rgb = SEVERITY_COLORS.get(report.overall_risk_rating, RGBColor(0, 0, 0))

    doc.add_heading("4. Findings Overview", level=1)
    with tempfile.TemporaryDirectory() as tmpdir:
        severity_path = f"{tmpdir}/severity.png"
        framework_path = f"{tmpdir}/framework.png"
        sev_result = render_severity_chart(report, severity_path)
        fw_result = render_framework_chart(report, framework_path)
        if sev_result:
            doc.add_picture(sev_result, width=Inches(6))
        if fw_result:
            doc.add_picture(fw_result, width=Inches(6))

    sorted_findings = sorted(
        report.findings,
        key=lambda f: SEVERITY_ORDER.index(f.severity) if f.severity in SEVERITY_ORDER else 99,
    )
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "#"
    header[1].text = "Title"
    header[2].text = "Severity"
    header[3].text = "Framework / Control"
    header[4].text = "Source"
    for row_cells in header:
        for paragraph in row_cells.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True

    if not sorted_findings:
        row = table.add_row()
        row.cells[1].text = "No findings"
    else:
        for i, finding in enumerate(sorted_findings, start=1):
            row = table.add_row()
            row.cells[0].text = str(i)
            row.cells[1].text = finding.title[:60] if len(finding.title) > 60 else finding.title
            sev_cell = row.cells[2]
            sev_cell.text = finding.severity
            for paragraph in sev_cell.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = SEVERITY_COLORS.get(finding.severity, RGBColor(0, 0, 0))
            if finding.framework and finding.control_id:
                row.cells[3].text = f"{finding.framework} — {finding.control_id}"
            else:
                row.cells[3].text = "—"
            row.cells[4].text = finding.source_file or "—"

    doc.add_heading("5. Findings", level=1)
    if not sorted_findings:
        doc.add_paragraph("No findings were identified in the reviewed material.")

    for i, finding in enumerate(sorted_findings, start=1):
        h = doc.add_heading(f"5.{i} {finding.title}", level=2)
        sev_p = doc.add_paragraph()
        sev_p.add_run("Severity: ").bold = True
        sev_run = sev_p.add_run(finding.severity)
        sev_run.bold = True
        sev_run.font.color.rgb = SEVERITY_COLORS.get(finding.severity, RGBColor(0, 0, 0))

        if finding.source_file:
            src_p = doc.add_paragraph()
            src_p.add_run("Source: ").bold = True
            src_p.add_run(finding.source_file)

        if finding.framework and finding.control_id:
            fw_p = doc.add_paragraph()
            fw_p.add_run("Framework Mapping: ").bold = True
            fw_p.add_run(f"{finding.framework} — {finding.control_id}")

        doc.add_paragraph("Description:", style="Intense Quote")
        doc.add_paragraph(finding.description)

        doc.add_paragraph("Evidence:", style="Intense Quote")
        doc.add_paragraph(finding.evidence)

        doc.add_paragraph("Remediation:", style="Intense Quote")
        doc.add_paragraph(finding.remediation)

        doc.add_paragraph()  # spacing

    doc.add_heading("6. Recommendations Summary", level=1)
    for rec in report.recommendations_summary:
        doc.add_paragraph(rec, style="List Bullet")

    if report.diagram and is_valid_mermaid(report.diagram):
        doc.add_page_break()
        doc.add_heading("7. Architecture Diagram (Mermaid source)", level=1)
        for line in report.diagram.splitlines():
            p = doc.add_paragraph(line)
            p.style = "No Spacing"
            for run in p.runs:
                run.font.name = "Courier New"
                run.font.size = Pt(9)

    doc.save(output_path)
    return output_path
