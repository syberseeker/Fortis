import os
import tempfile
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from .schema import SecurityReport, SEVERITY_ORDER
from .charts import render_severity_chart
from ..structured import is_valid_mermaid

SEVERITY_COLORS = {
    "Critical": RGBColor(0x8B, 0x00, 0x00),
    "High": RGBColor(0xC0, 0x39, 0x2B),
    "Medium": RGBColor(0xE6, 0x7E, 0x22),
    "Low": RGBColor(0x21, 0x87, 0x38),
    "Informational": RGBColor(0x5D, 0x6D, 0x7E),
}

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
SLATE = RGBColor(0x5D, 0x6D, 0x7E)


def _add_title_slide(prs, report: SecurityReport):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = report.title
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(36)
    subtitle = slide.placeholders[1]
    subtitle.text = f"Cybersecurity Assessment — {datetime.now().strftime('%d %B %Y')}"
    return slide


def _add_bullet_slide(prs, heading: str, bullets: list, sub=None):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = heading
    body = slide.placeholders[1].text_frame
    body.clear()
    first = True
    if sub:
        body.text = sub
        first = False
    for b in bullets:
        p = body.paragraphs[0] if first else body.add_paragraph()
        p.text = b
        p.level = 0
        first = False
    return slide


def _add_finding_slide(prs, idx: int, finding):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = f"Finding {idx}: {finding.title}"

    body = slide.placeholders[1].text_frame
    body.clear()

    p = body.paragraphs[0]
    run = p.add_run()
    run.text = f"Severity: {finding.severity}"
    run.font.bold = True
    run.font.color.rgb = SEVERITY_COLORS.get(finding.severity, RGBColor(0, 0, 0))

    if finding.source_file:
        p2 = body.add_paragraph()
        p2.text = f"Source: {finding.source_file}"

    if finding.framework and finding.control_id:
        p3 = body.add_paragraph()
        p3.text = f"Mapped control: {finding.framework} — {finding.control_id}"

    p4 = body.add_paragraph()
    p4.text = f"Description: {finding.description}"

    p5 = body.add_paragraph()
    p5.text = f"Remediation: {finding.remediation}"

    return slide


def render_pptx(report: SecurityReport, output_path: str) -> str:
    prs = Presentation()

    _add_title_slide(prs, report)

    _add_bullet_slide(
        prs,
        "Executive Summary",
        [report.executive_summary, f"Overall Risk Rating: {report.overall_risk_rating}"],
    )

    _add_bullet_slide(prs, "Scope & Context", [report.client_context, report.scope])

    sorted_findings = sorted(
        report.findings,
        key=lambda f: SEVERITY_ORDER.index(f.severity) if f.severity in SEVERITY_ORDER else 99,
    )

    severity_counts = {}
    for f in sorted_findings:
        severity_counts[f.severity] = severity_counts.get(f.severity, 0) + 1

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Findings Overview"

    severity_table_shape = slide.shapes.add_table(6, 2, Inches(0.5), Inches(1.5), Inches(9), Inches(2))
    severity_table = severity_table_shape.table
    severity_table.columns[0].width = Inches(3)
    severity_table.columns[1].width = Inches(2)

    header_cells = severity_table.rows[0].cells
    header_cells[0].text = "Severity"
    header_cells[1].text = "Count"
    for cell in header_cells:
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True

    for i, sev in enumerate(SEVERITY_ORDER, start=1):
        row = severity_table.rows[i]
        row.cells[0].text = sev
        row.cells[1].text = str(severity_counts.get(sev, 0))
        color = SEVERITY_COLORS.get(sev, RGBColor(0, 0, 0))
        for paragraph in row.cells[0].text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = color
                run.font.bold = True

    tmpdir = tempfile.mkdtemp()
    chart_path = None
    try:
        severity_path = os.path.join(tmpdir, "severity.png")
        sev_result = render_severity_chart(report, severity_path)
        if sev_result:
            slide.shapes.add_picture(sev_result, Inches(5), Inches(3.5), width=Inches(4))
            chart_path = sev_result

        if not sorted_findings:
            _add_bullet_slide(prs, "Findings", ["No findings identified in the reviewed material."])
        elif len(sorted_findings) <= 10:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = "Findings Summary"

            findings_table_shape = slide.shapes.add_table(len(sorted_findings) + 1, 4, Inches(0.5), Inches(1.5), Inches(9), Inches(5))
            findings_table = findings_table_shape.table
            findings_table.columns[0].width = Inches(0.5)
            findings_table.columns[1].width = Inches(4)
            findings_table.columns[2].width = Inches(1.5)
            findings_table.columns[3].width = Inches(3)

            header_cells = findings_table.rows[0].cells
            header_cells[0].text = "#"
            header_cells[1].text = "Finding"
            header_cells[2].text = "Severity"
            header_cells[3].text = "Framework"
            for cell in header_cells:
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True

            for i, finding in enumerate(sorted_findings, start=1):
                row = findings_table.rows[i]
                row.cells[0].text = str(i)
                row.cells[1].text = finding.title[:50] if len(finding.title) > 50 else finding.title
                row.cells[2].text = finding.severity
                color = SEVERITY_COLORS.get(finding.severity, RGBColor(0, 0, 0))
                for paragraph in row.cells[2].text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = color
                fw_text = f"{finding.framework} — {finding.control_id}" if finding.framework and finding.control_id else "—"
                row.cells[3].text = fw_text
        else:
            top_10 = sorted_findings[:10]
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = "Findings Summary (Top 10)"

            findings_table_shape = slide.shapes.add_table(11, 4, Inches(0.5), Inches(1.5), Inches(9), Inches(5))
            findings_table = findings_table_shape.table
            findings_table.columns[0].width = Inches(0.5)
            findings_table.columns[1].width = Inches(4)
            findings_table.columns[2].width = Inches(1.5)
            findings_table.columns[3].width = Inches(3)

            header_cells = findings_table.rows[0].cells
            header_cells[0].text = "#"
            header_cells[1].text = "Finding"
            header_cells[2].text = "Severity"
            header_cells[3].text = "Framework"
            for cell in header_cells:
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True

            for i, finding in enumerate(top_10, start=1):
                row = findings_table.rows[i]
                row.cells[0].text = str(i)
                row.cells[1].text = finding.title[:50] if len(finding.title) > 50 else finding.title
                row.cells[2].text = finding.severity
                color = SEVERITY_COLORS.get(finding.severity, RGBColor(0, 0, 0))
                for paragraph in row.cells[2].text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = color
                fw_text = f"{finding.framework} — {finding.control_id}" if finding.framework and finding.control_id else "—"
                row.cells[3].text = fw_text

        _add_bullet_slide(prs, "Recommendations", report.recommendations_summary)

        if report.diagram and is_valid_mermaid(report.diagram):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            title = slide.shapes.title
            title.text = "Architecture Diagram (source)"
            body = slide.placeholders[1].text_frame
            body.clear()
            body.word_wrap = True
            p = body.paragraphs[0]
            for i, line in enumerate(report.diagram.splitlines()):
                if i == 0:
                    p.text = line
                else:
                    p = body.add_paragraph()
                    p.text = line
                for run in p.runs:
                    run.font.name = "Courier New"
                    run.font.size = Pt(9)

        prs.save(output_path)
    finally:
        if chart_path:
            try:
                os.remove(chart_path)
            except OSError:
                pass
        try:
            os.rmdir(tmpdir)
        except OSError:
            pass

    return output_path
