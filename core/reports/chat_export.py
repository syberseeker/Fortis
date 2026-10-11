import re
from datetime import datetime

from markdown_it import MarkdownIt
from markdown_it.token import Token

_TAG_RE = re.compile(r"<[^>]+>")
_MAX_CODE_LINES = 25


def _plain_text(token: Token) -> str:
    parts = []
    for child in token.children or []:
        if child.type in ("text", "code_inline"):
            parts.append(child.content)
        elif child.type == "image":
            parts.append(_plain_text(child))
        elif child.type in ("softbreak", "hardbreak"):
            parts.append(" ")
    text = _TAG_RE.sub("", "".join(parts))
    return re.sub(r"\s+", " ", text).strip()


_MAX_TABLE_ROWS = 200
_MAX_TABLE_COLS = 12


def _table_block(rows):
    header, body = [], []
    for row in rows:
        if row and row[0][0] == "header":
            header.append([cell for _, cell in row])
        else:
            body.append([cell for _, cell in row])
    all_rows = header + body
    width = min(max((len(r) for r in all_rows), default=0), _MAX_TABLE_COLS)

    def _fit(cells):
        cells = cells[:_MAX_TABLE_COLS]
        return cells + ["…"] * max(0, width - len(cells))

    header_fit = _fit(header[0]) if header else []
    body_fit = [_fit(r) for r in body[:_MAX_TABLE_ROWS]]
    if len(body) > _MAX_TABLE_ROWS:
        body_fit.append(["…"] * width)
    return {"type": "table", "header": header_fit, "rows": body_fit}


def parse_reply_blocks(md_text: str):
    md = MarkdownIt("commonmark").enable(["table", "strikethrough"])
    tokens = md.parse(md_text)
    blocks = []
    quote_stack = []
    list_stack = []
    item_stack = []
    paragraph_mode = "text"
    current_list = None
    current_item = None
    table_rows = None
    current_row = None
    cell_kind = None
    cell_text = ""
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        ttype = tok.type
        if ttype == "heading_open":
            inline = tokens[i + 1] if i + 1 < len(tokens) and tokens[i + 1].type == "inline" else None
            blocks.append({"type": "heading", "level": int(tok.tag[1:]),
                           "text": _plain_text(inline) if inline else ""})
            i += 2
            continue
        if ttype == "paragraph_open":
            if current_item is None and not quote_stack:
                paragraph_mode = "text"
            inline = tokens[i + 1] if i + 1 < len(tokens) and tokens[i + 1].type == "inline" else None
            text = _plain_text(inline) if inline else ""
            if text:
                if paragraph_mode == "quote":
                    blocks.append({"type": "paragraph", "text": f"> {text}"})
                elif paragraph_mode == "item" and current_item is not None:
                    current_item.append(text)
                else:
                    blocks.append({"type": "paragraph", "text": text})
            i += 2
            continue
        if ttype == "blockquote_open":
            quote_stack.append("blockquote")
            paragraph_mode = "quote"
        elif ttype == "blockquote_close":
            if quote_stack:
                quote_stack.pop()
            paragraph_mode = "quote" if quote_stack else "text"
        elif ttype == "fence":
            blocks.append({"type": "code_block", "text": tok.content.rstrip("\n"), "language": tok.info.strip()})
        elif ttype == "code_block":
            blocks.append({"type": "code_block", "text": tok.content.rstrip("\n"), "language": ""})
        elif ttype == "html_block":
            text = re.sub(r"\s+", " ", _TAG_RE.sub("", tok.content)).strip()
            if text and paragraph_mode == "text":
                blocks.append({"type": "paragraph", "text": text})
        elif ttype == "hr":
            blocks.append({"type": "hr"})
        elif ttype in ("bullet_list_open", "ordered_list_open"):
            new_list = {
                "type": "bullet_list" if ttype == "bullet_list_open" else "ordered_list",
                "items": [],
            }
            blocks.append(new_list)
            if current_list is not None:
                list_stack.append(current_list)
            current_list = new_list
        elif ttype in ("bullet_list_close", "ordered_list_close"):
            current_list = list_stack.pop() if list_stack else None
        elif ttype == "list_item_open":
            current_item = []
            if current_list is not None:
                current_list["items"].append(current_item)
            item_stack.append(current_item)
            paragraph_mode = "item"
        elif ttype == "list_item_close":
            if item_stack:
                item_stack.pop()
            current_item = item_stack[-1] if item_stack else None
            if current_item is None:
                paragraph_mode = "quote" if quote_stack else "text"
        elif ttype == "table_open":
            table_rows = []
            current_row = None
        elif ttype == "tr_open" and table_rows is not None:
            current_row = []
            table_rows.append(current_row)
        elif ttype in ("th_open", "td_open") and table_rows is not None:
            cell_kind, cell_text = ttype[:2], ""
        elif ttype == "inline" and cell_kind is not None:
            cell_text = _plain_text(tok)
        elif ttype in ("th_close", "td_close") and table_rows is not None:
            if current_row is not None:
                current_row.append(("header" if cell_kind == "th" else "body", cell_text))
            cell_kind = None
        elif ttype == "table_close" and table_rows is not None:
            blocks.append(_table_block(table_rows))
            table_rows = None
            current_row = None
            cell_kind = None
        i += 1

    keep = []
    for block in blocks:
        if block["type"] in ("bullet_list", "ordered_list"):
            block["items"] = [" ".join(item).strip() for item in block["items"] if " ".join(item).strip()]
            if block["items"]:
                keep.append(block)
        elif block["type"] in ("heading", "code_block", "table", "hr"):
            keep.append(block)
        elif block["text"]:
            keep.append(block)
    return keep


def _export_date() -> str:
    return datetime.now().strftime("%d %B %Y")


def _require_text(md_text: str):
    if not md_text or not md_text.strip():
        raise ValueError("Nothing to export.")


def render_reply_markdown(md_text: str, output_path: str, title: str = "") -> str:
    _require_text(md_text)
    with open(output_path, "w", encoding="utf-8", newline="") as f:
        f.write(md_text)
    return output_path


def render_reply_docx(md_text: str, output_path: str, title: str = "") -> str:
    _require_text(md_text)
    from docx import Document
    from docx.shared import Pt

    blocks = parse_reply_blocks(md_text)
    doc = Document()

    if title:
        doc.add_heading(title, level=0)
        doc.add_paragraph(f"Chat export — {_export_date()}")
        if len(blocks) > 30:
            doc.add_page_break()

    for block in blocks:
        btype = block["type"]
        if btype == "heading":
            doc.add_heading(block["text"], level=min(block["level"] + 1, 4))
        elif btype == "paragraph":
            doc.add_paragraph(block["text"])
        elif btype == "bullet_list":
            for item in block["items"]:
                doc.add_paragraph(item, style="List Bullet")
        elif btype == "ordered_list":
            for item in block["items"]:
                doc.add_paragraph(item, style="List Number")
        elif btype == "code_block":
            for line in block["text"].splitlines() or [""]:
                p = doc.add_paragraph(line)
                p.style = "No Spacing"
                for run in p.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
        elif btype == "table":
            header = block["header"]
            rows = block["rows"]
            ncols = max([len(header)] + [len(r) for r in rows]) if header else max([len(r) for r in rows], default=0)
            nrows = (1 if header else 0) + len(rows)
            if ncols and nrows:
                table = doc.add_table(rows=nrows, cols=ncols)
                table.style = "Table Grid"
                ridx = 0
                if header:
                    for cidx in range(ncols):
                        table.rows[0].cells[cidx].text = header[cidx] if cidx < len(header) else ""
                    for cell in table.rows[0].cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.bold = True
                    ridx = 1
                for row in rows:
                    for cidx in range(ncols):
                        table.rows[ridx].cells[cidx].text = row[cidx] if cidx < len(row) else ""
                    ridx += 1
        elif btype == "hr":
            doc.add_paragraph("────────────────────────")

    doc.save(output_path)
    return output_path


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def render_reply_pdf(md_text: str, output_path: str, title: str = "") -> str:
    _require_text(md_text)
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem, Table, TableStyle
    )

    blocks = parse_reply_blocks(md_text)
    doc = SimpleDocTemplate(output_path, pagesize=LETTER,
                            topMargin=0.9 * inch, bottomMargin=0.9 * inch)
    ss = getSampleStyleSheet()
    ss.add(ParagraphStyle(name="CenterTitle", parent=ss["Title"], alignment=1))
    ss.add(ParagraphStyle(name="CenterSub", parent=ss["Normal"], alignment=1,
                          textColor=colors.HexColor("#5D6D7E"), fontSize=13))
    ss.add(ParagraphStyle(name="Mono", parent=ss["Code"], fontName="Courier", fontSize=9))

    story = []
    if title:
        story.append(Paragraph(_esc(title), ss["CenterTitle"]))
        story.append(Paragraph(f"Chat export — {_export_date()}", ss["CenterSub"]))
        story.append(Spacer(1, 0.3 * inch))

    for block in blocks:
        btype = block["type"]
        if btype == "heading":
            story.append(Paragraph(_esc(block["text"]), ss[f"Heading{min(block['level'], 4)}"]))
            story.append(Spacer(1, 0.08 * inch))
        elif btype == "paragraph":
            story.append(Paragraph(_esc(block["text"]), ss["BodyText"]))
            story.append(Spacer(1, 0.1 * inch))
        elif btype in ("bullet_list", "ordered_list"):
            items = [ListItem(Paragraph(_esc(item), ss["BodyText"])) for item in block["items"]]
            bullet_type = "bullet" if btype == "bullet_list" else "1"
            story.append(ListFlowable(items, bulletType=bullet_type))
            story.append(Spacer(1, 0.1 * inch))
        elif btype == "code_block":
            code = "<br/>".join(_esc(line) for line in block["text"].splitlines() or [""])
            story.append(Paragraph(f'<pre fontName="Courier" fontSize="9">{code}</pre>', ss["Mono"]))
            story.append(Spacer(1, 0.1 * inch))
        elif btype == "table":
            data = []
            if block["header"]:
                data.append([_esc(cell) for cell in block["header"]])
            for row in block["rows"]:
                data.append([_esc(cell) for cell in row])
            if data:
                ncols = max(len(r) for r in data)
                data = [r + [""] * (ncols - len(r)) for r in data]
                available = 6.7 * inch
                table = Table(data, colWidths=[available / ncols] * ncols)
                style = [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                ]
                if block["header"]:
                    style.append(("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey))
                    style.append(("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"))
                table.setStyle(TableStyle(style))
                story.append(table)
                story.append(Spacer(1, 0.15 * inch))
        elif btype == "hr":
            story.append(Spacer(1, 0.12 * inch))

    doc.build(story)
    return output_path


def render_reply_pptx(md_text: str, output_path: str, title: str = "") -> str:
    _require_text(md_text)
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu

    blocks = parse_reply_blocks(md_text)

    def first_text():
        for block in blocks:
            if block["type"] in ("heading", "paragraph") and block["text"]:
                return block["text"]
        return "Chat Export"

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = title or first_text()
    title_slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(36)
    title_slide.placeholders[1].text = f"Chat export — {_export_date()}"

    state = {"slide": None, "body": None, "used": False}

    def new_slide(heading):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = heading
        state["slide"] = slide
        state["body"] = None
        state["used"] = False

    def ensure_body():
        if state["body"] is None:
            if state["slide"] is None:
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = "Highlights"
                state["slide"] = slide
            try:
                body = state["slide"].placeholders[1].text_frame
            except KeyError:
                return None
            body.clear()
            body.word_wrap = True
            state["body"] = body
            state["used"] = False
        return state["body"]

    def add_para():
        tf = ensure_body()
        if tf is None:
            return None
        if state["used"]:
            p = tf.add_paragraph()
        else:
            p = tf.paragraphs[0]
            state["used"] = True
        return p

    def add_table(block):
        header, rows = block["header"], block["rows"]
        ncols = max([len(header)] + [len(r) for r in rows]) if header else max([len(r) for r in rows], default=0)
        nrows = (1 if header else 0) + len(rows)
        if not ncols or not nrows:
            return
        ensure_body()
        host = state["slide"]
        shape = host.shapes.add_table(nrows, ncols, Inches(0.5), Inches(1.7), Inches(9), Inches(0.4 * nrows + 0.2))
        tbl = shape.table
        weights = []
        for cidx in range(ncols):
            longest = 1
            cells = ([header[cidx]] if header and cidx < len(header) else []) + \
                    [r[cidx] if cidx < len(r) else "" for r in rows]
            for cell in cells:
                longest = max(longest, len(cell))
            weights.append(min(max(longest, 4), 40))
        total = Inches(9)
        used = 0
        for cidx in range(ncols):
            if cidx == ncols - 1:
                width = Emu(int(total) - used)
            else:
                width = Emu(int(int(total) * weights[cidx] / sum(weights)))
                used += int(width)
            tbl.columns[cidx].width = width
        ridx = 0
        if header:
            for cidx in range(ncols):
                tbl.rows[0].cells[cidx].text = header[cidx] if cidx < len(header) else ""
            for cell in tbl.rows[0].cells:
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
            ridx = 1
        for row in rows:
            for cidx in range(ncols):
                tbl.rows[ridx].cells[cidx].text = row[cidx] if cidx < len(row) else ""
            ridx += 1

    has_heading_slide = any(b["type"] == "heading" and b["level"] in (1, 2) for b in blocks)
    has_content = any(b["type"] in ("paragraph", "bullet_list", "ordered_list", "code_block", "table", "hr")
                      for b in blocks)

    for block in blocks:
        btype = block["type"]
        if btype == "heading":
            if block["level"] in (1, 2):
                new_slide(block["text"])
            else:
                p = add_para()
                if p is not None:
                    p.text = block["text"]
                    for run in p.runs:
                        run.font.bold = True
        elif btype == "paragraph":
            p = add_para()
            if p is not None:
                p.text = block["text"]
        elif btype in ("bullet_list", "ordered_list"):
            tf = ensure_body()
            if tf is not None:
                for n, item in enumerate(block["items"], start=1):
                    p = tf.add_paragraph() if state["used"] else tf.paragraphs[0]
                    p.text = f"{n}. {item}" if btype == "ordered_list" else item
                    p.level = 0
                    state["used"] = True
        elif btype == "code_block":
            lines = block["text"].splitlines()
            if len(lines) > _MAX_CODE_LINES:
                lines = lines[:_MAX_CODE_LINES] + ["..."]
            for line in lines or [""]:
                p = add_para()
                if p is None:
                    break
                p.text = line
                for run in p.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(11)
        elif btype == "table":
            add_table(block)
        elif btype == "hr":
            p = add_para()
            if p is not None:
                p.text = "————————————"

    if not has_heading_slide and has_content:
        ensure_body()

    prs.save(output_path)
    return output_path


REPLY_RENDERERS = {
    "md": render_reply_markdown,
    "docx": render_reply_docx,
    "pdf": render_reply_pdf,
    "pptx": render_reply_pptx,
}
